import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Load roster
wb_roster = openpyxl.load_workbook('รายชื่อ 1 สทค 2 (ม.6).xlsx', data_only=True)
s_roster = wb_roster['คะแนน']
ws_att = wb_roster['เช็กชื่อ']

students = []
att_counts = {}
for r in range(6, ws_att.max_row + 1):
    sid_val = ws_att.cell(r, 2).value
    if sid_val is not None:
        sid = str(int(sid_val)) if isinstance(sid_val, (int, float)) else str(sid_val).strip()
        cnt = int(ws_att.cell(r, 19).value)
        att_counts[sid] = cnt

# Manual overrides for attendance (จิตพิสัย)
att_counts['69319090021'] = 10  # น.ส.กาญจน์เกล้า แซ่ลิ้ม (เต็ม 20)
att_counts['69319090023'] = 9   # นายทักษดนย์ ชูชื่น (จิตพิสัย 18)
att_counts['69319090029'] = 9   # น.ส.ปิยธิดา บุตรนิน (จิตพิสัย 18)
att_counts['69319090031'] = 9   # นายพีรพัฒน์ พาหุรัตน์ (จิตพิสัย 18)
att_counts['69319090033'] = 9   # นายภานุวัฒน์ ยิ้มพ่วง (จิตพิสัย 18)
att_counts['69319090039'] = 9   # นายอนาวินทร์ แก้วเก้า (จิตพิสัย 18)

for r in range(6, s_roster.max_row + 1):
    seq_val = s_roster.cell(r, 1).value
    sid_val = s_roster.cell(r, 2).value
    name_val = s_roster.cell(r, 3).value
    if sid_val is not None:
        seq = int(seq_val) if seq_val is not None else len(students) + 1
        sid = str(int(sid_val)) if isinstance(sid_val, (int, float)) else str(sid_val).strip()
        name = str(name_val).strip()
        students.append({'seq': seq, 'sid': sid, 'name': name, 'sid3': sid[-3:]})

# Load classroom
wb_cr = openpyxl.load_workbook('คะแนน (31909-0003) การสร้างเว็บไซต์และระบบฐานข้อมูล 1 สทค 2 (พฤ. 8.00-13.00 น.xlsx', data_only=True)
s_cr = wb_cr['Sheet0']

student_cr_rows = {
    '69319090021': [12],
    '69319090022': [11],
    '69319090023': [14],
    '69319090024': [17],
    '69319090025': [10],
    '69319090026': [15],
    '69319090027': [8],
    '69319090028': [20],
    '69319090029': [9, 13, 16, 21],
    '69319090030': [18, 22],
    '69319090031': [7, 23],
    '69319090032': [24],
    '69319090033': [25],
    '69319090034': [26],
    '69319090035': [19],
    '69319090036': [6],
    '69319090037': [27],
    '69319090038': [29],
    '69319090039': [28],
}

asg_cols = {
    1: 10, 2: 33, 3: 34, 4: 35, 5: 20, 6: 23, 7: 26,
    8: 28, 9: 30, 10: 7, 11: 8, 12: 9, 13: 36
}

asg_headers = [
    ("Assignment 1\nBasic HTML", 10),
    ("Assignment 2\nBasic HTML 2", 10),
    ("Assignment 3\nBasic CSS", 10),
    ("Assignment 4\nColspan + Rowspan", 10),
    ("Assignment 5\nHTML Form", 10),
    ("Assignment 6\nCSS Boxmodel", 10),
    ("Assignment 7\nCSS Flexbox", 10),
    ("Assignment 8\nPHP Operator 1", 10),
    ("Assignment 9\nPHP Conditional", 10),
    ("Assignment 10\nArrays foreach Loop", 10),
    ("Assignment 11\nFunctions Modular", 10),
    ("Assignment 12\nForm Validation", 10),
    ("Assignment 13\nขึ้นระบบ Server", 10),
]

midterm_scores = {
    '021': (22.0, 22.0),
    '022': (17.0, 20.0),
    '023': (12.0, 5.0),
    '024': (27.0, 30.0),
    '025': (14.0, 5.0),
    '026': (18.0, 15.0),
    '027': (24.0, 28.0),
    '028': (17.0, 25.0),
    '029': (10.0, 15.0), # ปรับข้อเขียนกลางภาคให้ปิยธิดา (เดิม 0 เป็น 15)
    '030': (15.0, 13.0),
    '031': (17.0, 10.0), # ปรับข้อเขียนกลางภาคให้พีรพัฒน์ (เดิม 5 เป็น 10)
    '032': (20.0, 12.0),
    '033': (11.0, 12.0), # ปรับข้อเขียนกลางภาคให้ภานุวัฒน์ ยิ้มพ่วง (เดิม 0.5 เป็น 12)
    '034': (22.0, 25.0),
    '035': (12.0, 15.0),
    '036': (16.0, 17.0),
    '037': (12.0, 13.0),
    '038': (18.0, 19.0),
    '039': (11.0, 13.0),
}

final_scores = {
    '021': (24.0, 5.0),
    '022': (24.0, 30.0),
    '023': (17.0, 10.0), # ปรับข้อเขียนปลายภาคให้ทักษดนย์ (เดิม 0 เป็น 10)
    '024': (20.0, 13.0),
    '025': (17.0, 18.0),
    '026': (17.0, 33.0),
    '027': (19.0, 25.0),
    '028': (16.0, 8.0),
    '029': (12.0, 22.0), # ปรับข้อเขียนปลายภาคให้ปิยธิดา (เดิม 0 เป็น 22)
    '030': (15.0, 3.0),
    '031': (19.0, 15.0), # ปรับข้อเขียนปลายภาคให้พีรพัฒน์ (เดิม 0 เป็น 15)
    '032': (17.0, 25.0),
    '033': (11.0, 23.0), # ปรับข้อเขียนปลายภาคให้ภานุวัฒน์ ยิ้มพ่วง (เดิม 3 เป็น 23)
    '034': (18.0, 30.0),
    '035': (14.0, 2.0),
    '036': (20.0, 13.0),
    '037': (11.0, 23.5),
    '038': (16.0, 23.0),
    '039': (21.0, 18.0), # ปรับข้อเขียนปลายภาคให้อนาวินทร์ (เดิม 3 เป็น 18)
}

# Assignment overrides (ส่งงานย้อนหลัง)
asg_overrides = {
    '69319090021': {1: 10.0, 13: 10.0},
    '69319090023': {1: 10.0},
    '69319090029': {2: 10.0, 3: 10.0, 4: 10.0},
    '69319090031': {1: 10.0, 2: 10.0},
    '69319090033': {3: 10.0, 4: 10.0, 5: 10.0},
    '69319090039': {1: 10.0, 2: 10.0},
}

evidence_notes = {
    '69319090021': "ชื่อไฟล์แนบส่ง: 021w5.html (รหัสท้าย 021)",
    '69319090022': "ชื่อไฟล์แนบส่ง: 69319090022แบบฝึ...",
    '69319090023': "ชื่อบัญชีระบุ: ทักษดนย์ 023 (ตรงกับรหัสท้าย 023)",
    '69319090024': "ใช้อีเมลวิทยาลัย 69319090024@lbtech.ac.th, ไฟล์แนบ: 69319090024 week...",
    '69319090025': "ชื่อไฟล์แนบส่ง: 69319090025 input....",
    '69319090026': "ชื่อไฟล์แนบส่ง: 69319090026 งานแ...",
    '69319090027': "อีเมล nakarint1851@gmail.com (ชื่อ นครินทร์), ไฟล์แนบ: 69319090027 week...",
    '69319090028': "ชื่อไฟล์แนบส่ง: 69319090028 . .html",
    '69319090029': "พบใช้ 4 บัญชี: piyathdaboothnin@, preme5645@, ppzz123124@, 69319090029@ (คะแนนไม่ขัดแย้งกัน รวมเข้าด้วยกัน)",
    '69319090030': "พบ 2 บัญชี: 69319090030@lbtech.ac.th (มีคะแนน) และ pittawatwongchearw@ (0 คะแนน)",
    '69319090031': "พบ 2 บัญชี: skj19182@skj.ac.th (มีคะแนน, แนบไฟล์ 69319090031...) และ pccoco351@ (0 คะแนน)",
    '69319090032': "ชื่อบัญชีระบุ: ภาณุวัฒน์ ปั่นสันเที่ยะ 032, ไฟล์แนบ: 69319090032new.h...",
    '69319090033': "ชื่อบัญชีตรงกับรายชื่อ: นายภานุวัฒน์ ยิ้มพ่วง (สถานะส่งงานใน Classroom)",
    '69319090034': "ชื่อบัญชีระบุ: รชต เปลี่ยนศรี 070, ไฟล์แนบ: 69319090034.html",
    '69319090035': "ใช้อีเมลวิทยาลัย 69319090035@lbtech.ac.th และชื่อตรงกับรายชื่อ",
    '69319090036': "ชื่อบัญชีระบุ: 30ศิวา พุทธโชติ, ไฟล์แนบ: งานเว็ป69310036.html",
    '69319090037': "ชื่อตรงกับรายชื่อ ศุภกร แสงจันทร์, ไฟล์แนบ: งานเว็ป0037.html",
    '69319090038': "ชื่อบัญชีระบุ: โสธิตา มีผล 038 (ตรงกับรหัสท้าย 038)",
    '69319090039': "ชื่อบัญชีตรงกับรายชื่อ: นายอนาวินทร์ แก้วเก้า (สถานะส่งงานใน Classroom)",
}

# Create Workbook
wb_new = openpyxl.Workbook()

# Fonts
font_family = "Tahoma"
font_title = Font(name=font_family, size=14, bold=True, color="1B365D")
font_subtitle = Font(name=font_family, size=10, color="595959")
font_ratio = Font(name=font_family, size=10, bold=True, color="004D40")
font_super_header = Font(name=font_family, size=10, bold=True, color="FFFFFF")
font_header = Font(name=font_family, size=9, bold=True, color="FFFFFF")
font_data = Font(name=font_family, size=9.5, color="000000")
font_bold = Font(name=font_family, size=9.5, bold=True, color="000000")
font_stat_title = Font(name=font_family, size=9.5, bold=True, color="1B365D")
font_stat_val = Font(name=font_family, size=9.5, bold=True, color="1B365D")

# Fills
fill_hdr_main = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
fill_hdr_att = PatternFill(start_color="D35400", end_color="D35400", fill_type="solid")
fill_hdr_asg = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
fill_hdr_mid = PatternFill(start_color="1E8449", end_color="1E8449", fill_type="solid")
fill_hdr_fin = PatternFill(start_color="6C3483", end_color="6C3483", fill_type="solid")
fill_hdr_tot = PatternFill(start_color="0E6251", end_color="0E6251", fill_type="solid")
fill_hdr_grd = PatternFill(start_color="283747", end_color="283747", fill_type="solid")

fill_zebra = PatternFill(start_color="F8F9FA", end_color="F8F9FA", fill_type="solid")
fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
fill_stat = PatternFill(start_color="EAEDED", end_color="EAEDED", fill_type="solid")
fill_zero = PatternFill(start_color="FADBD8", end_color="FADBD8", fill_type="solid")

thin_border_side = Side(border_style="thin", color="D5D8DC")
thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
header_border = Border(left=Side(border_style="thin", color="FFFFFF"),
                       right=Side(border_style="thin", color="FFFFFF"),
                       top=Side(border_style="thin", color="FFFFFF"),
                       bottom=Side(border_style="medium", color="1B365D"))
double_bottom_border = Border(left=thin_border_side, right=thin_border_side,
                              top=thin_border_side, bottom=Side(border_style="double", color="1B365D"))

align_center = Alignment(horizontal="center", vertical="center")
align_left = Alignment(horizontal="left", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")
align_header = Alignment(horizontal="center", vertical="center", wrap_text=True)


# ====================================================
# SHEET 1: สรุปผลการเรียนและตัดเกรด (Executive Summary)
# ====================================================
ws1 = wb_new.active
ws1.title = "สรุปผลการเรียนและตัดเกรด"
ws1.views.sheetView[0].showGridLines = True

ws1.merge_cells("A1:N1")
ws1["A1"] = "รายงานสรุปผลการประเมินและตัดเกรดประจำวิชา การสร้างเว็บไซต์และระบบฐานข้อมูล (31909-0003)"
ws1["A1"].font = font_title; ws1["A1"].alignment = align_left

ws1.merge_cells("A2:N2")
ws1["A2"] = "ระดับชั้น ปวส.1 สาขาวิชาเทคโนโลยีคอมพิวเตอร์ (กลุ่ม 1 สทค 2 (ม.6)) | ภาคเรียนที่ 1/2569 | ครูผู้สอน: นายไพบูลย์ สมนึก"
ws1["A2"].font = font_subtitle; ws1["A2"].alignment = align_left

ws1.merge_cells("A3:N3")
ws1["A3"] = "เกณฑ์การตัดเกรด: จิตพิสัย (เวลาเรียน) 20% : คะแนนเก็บ 50% : สอบกลางภาค 10% : สอบปลายภาค 20% (รวม 100%)"
ws1["A3"].font = font_ratio; ws1["A3"].alignment = align_left

ws1.row_dimensions[1].height = 24
ws1.row_dimensions[2].height = 18
ws1.row_dimensions[3].height = 18
ws1.row_dimensions[4].height = 8
ws1.row_dimensions[5].height = 24
ws1.row_dimensions[6].height = 36

super_hdrs_s1 = [
    ("A5:C5", "ข้อมูลนักศึกษา", fill_hdr_main),
    ("D5:E5", "จิตพิสัย (เวลาเรียน 20%)", fill_hdr_att),
    ("F5:G5", "คะแนนเก็บ (50%)", fill_hdr_asg),
    ("H5:I5", "สอบกลางภาค (10%)", fill_hdr_mid),
    ("J5:K5", "สอบปลายภาค (20%)", fill_hdr_fin),
    ("L5:N5", "สรุปผลการเรียน (100%)", fill_hdr_tot),
]

for rng, txt, f_fill in super_hdrs_s1:
    ws1.merge_cells(rng)
    top_left = rng.split(":")[0]
    ws1[top_left] = txt
    ws1[top_left].font = font_super_header
    ws1[top_left].fill = f_fill
    ws1[top_left].alignment = align_header
    cells_in_rng = ws1[rng]
    for row in cells_in_rng:
        for cell in row:
            cell.fill = f_fill
            cell.border = header_border

sub_hdrs_s1 = [
    ("ลำดับที่", fill_hdr_main),
    ("รหัสนักศึกษา", fill_hdr_main),
    ("ชื่อ - สกุล", fill_hdr_main),
    ("มาเรียน\n(10 ครั้ง)", fill_hdr_att),
    ("จิตพิสัย\n(เต็ม 20)", fill_hdr_att),
    ("คะแนนดิบ\n(เต็ม 130)", fill_hdr_asg),
    ("ทอนคะแนน\n(เต็ม 50)", fill_hdr_asg),
    ("คะแนนดิบ\n(เต็ม 100)", fill_hdr_mid),
    ("ทอนคะแนน\n(เต็ม 10)", fill_hdr_mid),
    ("คะแนนดิบ\n(เต็ม 100)", fill_hdr_fin),
    ("ทอนคะแนน\n(เต็ม 20)", fill_hdr_fin),
    ("คะแนนรวม\n(เต็ม 100)", fill_hdr_tot),
    ("ระดับเกรด\n(0 - 4)", fill_hdr_grd),
    ("ผลการประเมิน", fill_hdr_grd),
]

for col_idx, (h_title, h_fill) in enumerate(sub_hdrs_s1, start=1):
    cell = ws1.cell(6, col_idx, h_title)
    cell.font = font_header
    cell.fill = h_fill
    cell.alignment = align_header
    cell.border = header_border

start_row = 7
for s_idx, s in enumerate(students):
    curr_row = start_row + s_idx
    raw_sheet_row = 6 + s_idx
    ws1.row_dimensions[curr_row].height = 20
    is_even = (s_idx % 2 == 1)
    row_fill = fill_zebra if is_even else fill_white

    cA = ws1.cell(curr_row, 1, s['seq']); cA.font = font_data; cA.alignment = align_center; cA.fill = row_fill; cA.border = thin_border
    cB = ws1.cell(curr_row, 2, s['sid']); cB.font = font_data; cB.alignment = align_center; cB.fill = row_fill; cB.border = thin_border; cB.number_format = '@'
    cC = ws1.cell(curr_row, 3, s['name']); cC.font = font_data; cC.alignment = align_left; cC.fill = row_fill; cC.border = thin_border

    # Attendance
    att_cnt = att_counts.get(s['sid'], 0)
    cD = ws1.cell(curr_row, 4, att_cnt); cD.font = font_data; cD.alignment = align_center; cD.fill = row_fill; cD.border = thin_border; cD.number_format = '0'
    cE = ws1.cell(curr_row, 5, f"=ROUND((D{curr_row}/10)*20, 2)"); cE.font = font_bold; cE.alignment = align_right; cE.fill = row_fill; cE.border = thin_border; cE.number_format = '0.00'

    # Assignment
    cF = ws1.cell(curr_row, 6, f"='คะแนนดิบ'!Q{raw_sheet_row}"); cF.font = font_data; cF.alignment = align_right; cF.fill = row_fill; cF.border = thin_border; cF.number_format = '0.0'
    cG = ws1.cell(curr_row, 7, f"=ROUND((F{curr_row}/130)*50, 2)"); cG.font = font_bold; cG.alignment = align_right; cG.fill = row_fill; cG.border = thin_border; cG.number_format = '0.00'

    # Midterm
    cH = ws1.cell(curr_row, 8, f"='คะแนนดิบ'!U{raw_sheet_row}"); cH.font = font_data; cH.alignment = align_right; cH.fill = row_fill; cH.border = thin_border; cH.number_format = '0.0'
    cI = ws1.cell(curr_row, 9, f"=ROUND((H{curr_row}/100)*10, 2)"); cI.font = font_bold; cI.alignment = align_right; cI.fill = row_fill; cI.border = thin_border; cI.number_format = '0.00'

    # Final
    cJ = ws1.cell(curr_row, 10, f"='คะแนนดิบ'!X{raw_sheet_row}"); cJ.font = font_data; cJ.alignment = align_right; cJ.fill = row_fill; cJ.border = thin_border; cJ.number_format = '0.0'
    cK = ws1.cell(curr_row, 11, f"=ROUND((J{curr_row}/100)*20, 2)"); cK.font = font_bold; cK.alignment = align_right; cK.fill = row_fill; cK.border = thin_border; cK.number_format = '0.00'

    # Total Net Score
    cL = ws1.cell(curr_row, 12, f"=ROUND(E{curr_row}+G{curr_row}+I{curr_row}+K{curr_row}, 2)")
    cL.font = font_bold; cL.alignment = align_right; cL.fill = row_fill; cL.border = thin_border; cL.number_format = '0.00'

    # Grade
    cM = ws1.cell(curr_row, 13, f"=IF(L{curr_row}>=80, 4, IF(L{curr_row}>=75, 3.5, IF(L{curr_row}>=70, 3, IF(L{curr_row}>=65, 2.5, IF(L{curr_row}>=60, 2, IF(L{curr_row}>=55, 1.5, IF(L{curr_row}>=50, 1, 0)))))))")
    cM.font = font_bold; cM.alignment = align_center; cM.fill = row_fill; cM.border = thin_border; cM.number_format = '0.0'

    # Pass/Fail
    cN = ws1.cell(curr_row, 14, f'=IF(M{curr_row}>0, "ผ่าน", "ไม่ผ่าน")')
    cN.font = font_bold; cN.alignment = align_center; cN.fill = row_fill; cN.border = thin_border

end_row_s1 = start_row + len(students) - 1

stat_rows_s1 = [
    ("คะแนนเฉลี่ย", "AVERAGE", thin_border),
    ("คะแนนสูงสุด", "MAX", thin_border),
    ("คะแนนต่ำสุด", "MIN", double_bottom_border)
]

for stat_idx, (label, func, b_style) in enumerate(stat_rows_s1):
    stat_r = end_row_s1 + 1 + stat_idx
    ws1.row_dimensions[stat_r].height = 21
    ws1.merge_cells(start_row=stat_r, start_column=1, end_row=stat_r, end_column=3)
    c_label = ws1.cell(stat_r, 1, label)
    c_label.font = font_stat_title; c_label.alignment = align_right; c_label.fill = fill_stat
    for c in range(1, 4):
        ws1.cell(stat_r, c).border = b_style; ws1.cell(stat_r, c).fill = fill_stat

    for c in range(4, 13):
        col_letter = get_column_letter(c)
        c_stat = ws1.cell(stat_r, c, f"={func}({col_letter}{start_row}:{col_letter}{end_row_s1})")
        c_stat.font = font_stat_val; c_stat.alignment = align_right; c_stat.fill = fill_stat; c_stat.border = b_style
        c_stat.number_format = '0.00' if func == 'AVERAGE' else ('0' if c == 4 else '0.00')

    ws1.cell(stat_r, 13).border = b_style; ws1.cell(stat_r, 13).fill = fill_stat
    ws1.cell(stat_r, 14).border = b_style; ws1.cell(stat_r, 14).fill = fill_stat

# Grade Distribution Box
box_start_r = 31
ws1.merge_cells(f"C{box_start_r}:F{box_start_r}")
ws1[f"C{box_start_r}"] = "สรุปการกระจายระดับผลการเรียน (Grade Distribution)"
ws1[f"C{box_start_r}"].font = font_super_header
ws1[f"C{box_start_r}"].fill = fill_hdr_main
ws1[f"C{box_start_r}"].alignment = align_header
for c in range(3, 7): ws1.cell(box_start_r, c).fill = fill_hdr_main; ws1.cell(box_start_r, c).border = header_border

dist_headers = [("ระดับเกรด", 3), ("ช่วงคะแนน", 4), ("จำนวน (คน)", 5), ("ร้อยละ (%)", 6)]
for title, col_idx in dist_headers:
    cell = ws1.cell(box_start_r + 1, col_idx, title)
    cell.font = font_header; cell.fill = fill_hdr_tot; cell.alignment = align_header; cell.border = header_border

grade_ranges = [
    (4.0, "80.00 - 100.00"),
    (3.5, "75.00 - 79.99"),
    (3.0, "70.00 - 74.99"),
    (2.5, "65.00 - 69.99"),
    (2.0, "60.00 - 64.99"),
    (1.5, "55.00 - 59.99"),
    (1.0, "50.00 - 54.99"),
    (0.0, "ต่ำกว่า 50.00"),
]

for g_idx, (grd_val, rng_txt) in enumerate(grade_ranges):
    r_idx = box_start_r + 2 + g_idx
    ws1.row_dimensions[r_idx].height = 19
    r_fill = fill_zebra if g_idx % 2 == 1 else fill_white
    
    cG = ws1.cell(r_idx, 3, grd_val); cG.font = font_bold; cG.alignment = align_center; cG.fill = r_fill; cG.border = thin_border; cG.number_format = '0.0'
    cR = ws1.cell(r_idx, 4, rng_txt); cR.font = font_data; cR.alignment = align_center; cR.fill = r_fill; cR.border = thin_border
    cC = ws1.cell(r_idx, 5, f'=COUNTIF(M${start_row}:M${end_row_s1}, {grd_val})'); cC.font = font_bold; cC.alignment = align_center; cC.fill = r_fill; cC.border = thin_border
    cP = ws1.cell(r_idx, 6, f'=ROUND((E{r_idx}/{len(students)})*100, 2)'); cP.font = font_bold; cP.alignment = align_right; cP.fill = r_fill; cP.border = thin_border; cP.number_format = '0.00'

tot_r_idx = box_start_r + 2 + len(grade_ranges)
ws1.merge_cells(f"C{tot_r_idx}:D{tot_r_idx}")
c_tot_lbl = ws1.cell(tot_r_idx, 3, "รวมทั้งหมด")
c_tot_lbl.font = font_bold; c_tot_lbl.alignment = align_right; c_tot_lbl.fill = fill_stat
ws1.cell(tot_r_idx, 4).fill = fill_stat
ws1.cell(tot_r_idx, 3).border = double_bottom_border; ws1.cell(tot_r_idx, 4).border = double_bottom_border

c_tot_cnt = ws1.cell(tot_r_idx, 5, f'=SUM(E{box_start_r+2}:E{tot_r_idx-1})')
c_tot_cnt.font = font_bold; c_tot_cnt.alignment = align_center; c_tot_cnt.fill = fill_stat; c_tot_cnt.border = double_bottom_border

c_tot_pct = ws1.cell(tot_r_idx, 6, f'=SUM(F{box_start_r+2}:F{tot_r_idx-1})')
c_tot_pct.font = font_bold; c_tot_pct.alignment = align_right; c_tot_pct.fill = fill_stat; c_tot_pct.border = double_bottom_border; c_tot_pct.number_format = '0.00'

ws1.freeze_panes = "D7"

ws1.column_dimensions['A'].width = 8
ws1.column_dimensions['B'].width = 16
ws1.column_dimensions['C'].width = 28
ws1.column_dimensions['D'].width = 12
ws1.column_dimensions['E'].width = 13
ws1.column_dimensions['F'].width = 13
ws1.column_dimensions['G'].width = 13
ws1.column_dimensions['H'].width = 13
ws1.column_dimensions['I'].width = 13
ws1.column_dimensions['J'].width = 13
ws1.column_dimensions['K'].width = 13
ws1.column_dimensions['L'].width = 15
ws1.column_dimensions['M'].width = 13
ws1.column_dimensions['N'].width = 14


# ====================================================
# SHEET 2: คะแนนดิบ (Assignments + Midterm + Final)
# ====================================================
ws2 = wb_new.create_sheet(title="คะแนนดิบ")
ws2.views.sheetView[0].showGridLines = True

ws2.merge_cells("A1:X1")
ws2["A1"] = "ตารางคะแนนดิบ: คะแนนเก็บตามภาระงาน และคะแนนสอบกลางภาค/ปลายภาค"
ws2["A1"].font = font_title; ws2["A1"].alignment = align_left

ws2.merge_cells("A2:X2")
ws2["A2"] = "วิชา การสร้างเว็บไซต์และระบบฐานข้อมูล (31909-0003) | ปวส.1 เทคโนโลยีคอมพิวเตอร์ (กลุ่ม 1 สทค 2 (ม.6)) | ภาคเรียนที่ 1/2569"
ws2["A2"].font = font_subtitle; ws2["A2"].alignment = align_left

ws2.row_dimensions[1].height = 24
ws2.row_dimensions[2].height = 18
ws2.row_dimensions[3].height = 8
ws2.row_dimensions[4].height = 24
ws2.row_dimensions[5].height = 38

ws2.merge_cells("A4:C4")
ws2["A4"] = "ข้อมูลนักศึกษา"
ws2["A4"].font = font_super_header; ws2["A4"].fill = fill_hdr_main; ws2["A4"].alignment = align_header
for c in range(1, 4): ws2.cell(4, c).border = header_border; ws2.cell(4, c).fill = fill_hdr_main

ws2.merge_cells("D4:R4")
ws2["D4"] = "คะแนนเก็บตามภาระงาน (Assignment 1 - 13 | เต็ม 130 คะแนน)"
ws2["D4"].font = font_super_header; ws2["D4"].fill = fill_hdr_asg; ws2["D4"].alignment = align_header
for c in range(4, 19): ws2.cell(4, c).border = header_border; ws2.cell(4, c).fill = fill_hdr_asg

ws2.merge_cells("S4:U4")
ws2["S4"] = "สอบกลางภาค (เต็ม 100 คะแนน)"
ws2["S4"].font = font_super_header; ws2["S4"].fill = fill_hdr_mid; ws2["S4"].alignment = align_header
for c in range(19, 22): ws2.cell(4, c).border = header_border; ws2.cell(4, c).fill = fill_hdr_mid

ws2.merge_cells("V4:X4")
ws2["V4"] = "สอบปลายภาค (เต็ม 100 คะแนน)"
ws2["V4"].font = font_super_header; ws2["V4"].fill = fill_hdr_fin; ws2["V4"].alignment = align_header
for c in range(22, 25): ws2.cell(4, c).border = header_border; ws2.cell(4, c).fill = fill_hdr_fin

sub_headers_s2 = [
    ("ลำดับที่", fill_hdr_main),
    ("รหัสนักศึกษา", fill_hdr_main),
    ("ชื่อ - สกุล", fill_hdr_main),
]
for title, max_p in asg_headers:
    sub_headers_s2.append((f"{title}\n({max_p})", fill_hdr_asg))
sub_headers_s2.append(("รวมงาน\n(130)", fill_hdr_main))
sub_headers_s2.append(("ร้อยละ\n(%)", fill_hdr_main))

sub_headers_s2.append(("ข้อกา\nปรนัย (40)", fill_hdr_mid))
sub_headers_s2.append(("ข้อเขียน\nอัตนัย (60)", fill_hdr_mid))
sub_headers_s2.append(("รวมกลางภาค\n(100)", fill_hdr_main))

sub_headers_s2.append(("ข้อกา\nปรนัย (40)", fill_hdr_fin))
sub_headers_s2.append(("ข้อเขียน\nอัตนัย (60)", fill_hdr_fin))
sub_headers_s2.append(("รวมปลายภาค\n(100)", fill_hdr_main))

for col_idx, (h_title, h_fill) in enumerate(sub_headers_s2, start=1):
    cell = ws2.cell(5, col_idx, h_title)
    cell.font = font_header; cell.fill = h_fill; cell.alignment = align_header; cell.border = header_border

start_row_s2 = 6
for s_idx, s in enumerate(students):
    curr_row = start_row_s2 + s_idx
    ws2.row_dimensions[curr_row].height = 20
    is_even = (s_idx % 2 == 1)
    row_fill = fill_zebra if is_even else fill_white

    cA = ws2.cell(curr_row, 1, s['seq']); cA.font = font_data; cA.alignment = align_center; cA.fill = row_fill; cA.border = thin_border
    cB = ws2.cell(curr_row, 2, s['sid']); cB.font = font_data; cB.alignment = align_center; cB.fill = row_fill; cB.border = thin_border; cB.number_format = '@'
    cC = ws2.cell(curr_row, 3, s['name']); cC.font = font_data; cC.alignment = align_left; cC.fill = row_fill; cC.border = thin_border

    cr_rows = student_cr_rows.get(s['sid'], [])
    scores_by_asg = {}
    for cr_r in cr_rows:
        for asg_num, cr_col in asg_cols.items():
            val = s_cr.cell(cr_r, cr_col).value
            if val is not None:
                if asg_num not in scores_by_asg or val > scores_by_asg[asg_num]:
                    scores_by_asg[asg_num] = float(val)

    # Apply assignment overrides
    if s['sid'] in asg_overrides:
        for asg_k, asg_v in asg_overrides[s['sid']].items():
            scores_by_asg[asg_k] = asg_v

    for asg_num in range(1, 14):
        col_idx = 3 + asg_num
        val = scores_by_asg.get(asg_num, 0.0)
        cell = ws2.cell(curr_row, col_idx, val)
        cell.font = font_data; cell.alignment = align_center; cell.border = thin_border
        cell.number_format = '0.0' if val % 1 != 0 else '0'
        if val == 0.0:
            cell.fill = fill_zero
        else:
            cell.fill = row_fill

    # Col Q: รวมงาน (130)
    col_Q = ws2.cell(curr_row, 17, f"=SUM(D{curr_row}:P{curr_row})")
    col_Q.font = font_bold; col_Q.alignment = align_right; col_Q.fill = row_fill; col_Q.border = thin_border; col_Q.number_format = '0.0'

    # Col R: ร้อยละ (%)
    col_R = ws2.cell(curr_row, 18, f"=ROUND((Q{curr_row}/130)*100, 2)")
    col_R.font = font_bold; col_R.alignment = align_right; col_R.fill = row_fill; col_R.border = thin_border; col_R.number_format = '0.00'

    # Midterm
    m_obj, m_subj = midterm_scores.get(s['sid3'], (0.0, 0.0))
    cS = ws2.cell(curr_row, 19, m_obj); cS.font = font_data; cS.alignment = align_center; cS.fill = row_fill; cS.border = thin_border; cS.number_format = '0.0' if m_obj % 1 != 0 else '0'
    cT = ws2.cell(curr_row, 20, m_subj); cT.font = font_data; cT.alignment = align_center; cT.fill = row_fill; cT.border = thin_border; cT.number_format = '0.0' if m_subj % 1 != 0 else '0'
    cU = ws2.cell(curr_row, 21, f"=SUM(S{curr_row}:T{curr_row})"); cU.font = font_bold; cU.alignment = align_right; cU.fill = row_fill; cU.border = thin_border; cU.number_format = '0.0'

    # Final
    f_obj, f_subj = final_scores.get(s['sid3'], (0.0, 0.0))
    cV = ws2.cell(curr_row, 22, f_obj); cV.font = font_data; cV.alignment = align_center; cV.fill = row_fill; cV.border = thin_border; cV.number_format = '0.0' if f_obj % 1 != 0 else '0'
    cW = ws2.cell(curr_row, 23, f_subj); cW.font = font_data; cW.alignment = align_center; cW.fill = row_fill; cW.border = thin_border; cW.number_format = '0.0' if f_subj % 1 != 0 else '0'
    cX = ws2.cell(curr_row, 24, f"=SUM(V{curr_row}:W{curr_row})"); cX.font = font_bold; cX.alignment = align_right; cX.fill = row_fill; cX.border = thin_border; cX.number_format = '0.0'

end_row_s2 = start_row_s2 + len(students) - 1

for stat_idx, (label, func, b_style) in enumerate(stat_rows_s1):
    stat_r = end_row_s2 + 1 + stat_idx
    ws2.row_dimensions[stat_r].height = 21
    ws2.merge_cells(start_row=stat_r, start_column=1, end_row=stat_r, end_column=3)
    c_label = ws2.cell(stat_r, 1, label)
    c_label.font = font_stat_title; c_label.alignment = align_right; c_label.fill = fill_stat
    for c in range(1, 4):
        ws2.cell(stat_r, c).border = b_style; ws2.cell(stat_r, c).fill = fill_stat

    for c in range(4, 25):
        col_letter = get_column_letter(c)
        c_stat = ws2.cell(stat_r, c, f"={func}({col_letter}{start_row_s2}:{col_letter}{end_row_s2})")
        c_stat.font = font_stat_val; c_stat.alignment = align_right if c in [17, 18, 21, 24] else align_center
        c_stat.fill = fill_stat; c_stat.border = b_style
        c_stat.number_format = '0.00' if (func == 'AVERAGE' or c == 18) else '0.0'

ws2.freeze_panes = "D6"

ws2.column_dimensions['A'].width = 8
ws2.column_dimensions['B'].width = 16
ws2.column_dimensions['C'].width = 28
for c in range(4, 17):
    ws2.column_dimensions[get_column_letter(c)].width = 15
ws2.column_dimensions['Q'].width = 14
ws2.column_dimensions['R'].width = 13
ws2.column_dimensions['S'].width = 14
ws2.column_dimensions['T'].width = 14
ws2.column_dimensions['U'].width = 15
ws2.column_dimensions['V'].width = 14
ws2.column_dimensions['W'].width = 14
ws2.column_dimensions['X'].width = 15


# ====================================================
# SHEET 3: จับคู่บัญชี-อีเมล
# ====================================================
ws3 = wb_new.create_sheet(title="จับคู่บัญชี-อีเมล")
ws3.views.sheetView[0].showGridLines = True

ws3.merge_cells("A1:G1")
ws3["A1"] = "ตารางจับคู่บัญชี Google Classroom กับทะเบียนนักศึกษา"
ws3["A1"].font = font_title; ws3["A1"].alignment = align_left

ws3.merge_cells("A2:G2")
ws3["A2"] = "วิชา การสร้างเว็บไซต์และระบบฐานข้อมูล (31909-0003) | 1 สทค 2 (ม.6) | ใช้ตรวจสอบความถูกต้องของการลงคะแนน"
ws3["A2"].font = font_subtitle; ws3["A2"].alignment = align_left

ws3.row_dimensions[1].height = 25
ws3.row_dimensions[2].height = 18
ws3.row_dimensions[3].height = 8
ws3.row_dimensions[4].height = 28

headers_m3 = [
    "ลำดับที่",
    "รหัสนักศึกษา",
    "ชื่อ - สกุล ในทะเบียน",
    "ชื่อที่แสดงใน Classroom",
    "อีเมลที่ใช้ส่งงานใน Classroom",
    "จำนวนงานที่พบ",
    "หลักฐานการจับคู่ / บัญชีเพิ่มเติม"
]

for col_idx, h_title in enumerate(headers_m3, start=1):
    cell = ws3.cell(4, col_idx, h_title)
    cell.font = font_header; cell.fill = fill_hdr_main; cell.alignment = align_header; cell.border = header_border

for s_idx, s in enumerate(students):
    curr_row = 5 + s_idx
    ws3.row_dimensions[curr_row].height = 22
    is_even = (s_idx % 2 == 1)
    row_fill = fill_zebra if is_even else fill_white

    cr_rows = student_cr_rows.get(s['sid'], [])
    cr_names = []
    cr_emails = []
    for cr_r in cr_rows:
        f = str(s_cr.cell(cr_r, 2).value or '').strip()
        l = str(s_cr.cell(cr_r, 1).value or '').strip()
        em = str(s_cr.cell(cr_r, 3).value or '').strip()
        full_name = f"{f} {l}".strip()
        if full_name: cr_names.append(full_name)
        if em: cr_emails.append(em)

    scores_by_asg = {}
    for cr_r in cr_rows:
        for asg_num, cr_col in asg_cols.items():
            val = s_cr.cell(cr_r, cr_col).value
            if val is not None:
                scores_by_asg[asg_num] = val

    if s['sid'] in asg_overrides:
        for asg_k, asg_v in asg_overrides[s['sid']].items():
            scores_by_asg[asg_k] = asg_v

    cA = ws3.cell(curr_row, 1, s['seq']); cA.font = font_data; cA.alignment = align_center; cA.fill = row_fill; cA.border = thin_border
    cB = ws3.cell(curr_row, 2, s['sid']); cB.font = font_data; cB.alignment = align_center; cB.fill = row_fill; cB.border = thin_border; cB.number_format = '@'
    cC = ws3.cell(curr_row, 3, s['name']); cC.font = font_data; cC.alignment = align_left; cC.fill = row_fill; cC.border = thin_border
    cD = ws3.cell(curr_row, 4, " / ".join(cr_names)); cD.font = font_data; cD.alignment = align_left; cD.fill = row_fill; cD.border = thin_border
    cE = ws3.cell(curr_row, 5, ", ".join(cr_emails)); cE.font = font_data; cE.alignment = align_left; cE.fill = row_fill; cE.border = thin_border
    cF = ws3.cell(curr_row, 6, f"{len(scores_by_asg)} / 13 งาน"); cF.font = font_bold; cF.alignment = align_center; cF.fill = row_fill; cF.border = thin_border
    cG = ws3.cell(curr_row, 7, evidence_notes.get(s['sid'], "")); cG.font = font_data; cG.alignment = align_left; cG.fill = row_fill; cG.border = thin_border

ws3.freeze_panes = "D5"

ws3.column_dimensions['A'].width = 8
ws3.column_dimensions['B'].width = 16
ws3.column_dimensions['C'].width = 28
ws3.column_dimensions['D'].width = 30
ws3.column_dimensions['E'].width = 45
ws3.column_dimensions['F'].width = 16
ws3.column_dimensions['G'].width = 55

output_path = 'สรุปคะแนน_1สทค2.xlsx'
wb_new.save(output_path)
print(f"Successfully generated complete graded workbook '{output_path}'!")
